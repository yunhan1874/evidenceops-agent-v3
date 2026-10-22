#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "NEW_VERSION_FEATURES_AND_COMPARISON_2026-07-26.md"
DEFAULT_OUTPUT = ROOT / "NEW_VERSION_FEATURES_AND_COMPARISON_2026-07-26.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_in: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    total_dxa = sum(int(width * 1440) for width in widths_in)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_in:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(grid_col)

    for row in table.rows:
        for index, width in enumerate(widths_in):
            if index >= len(row.cells):
                continue
            cell = row.cells[index]
            cell.width = Inches(width)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def style_run(run, bold=False, italic=False, code=False) -> None:
    run.bold = bold
    run.italic = italic
    if code:
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(9.5)


def add_markdown_runs(paragraph, text: str) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            style_run(run, code=True)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            style_run(run, bold=True)
        else:
            paragraph.add_run(part)


def add_paragraph(doc: Document, text: str, style: str = "Normal"):
    paragraph = doc.add_paragraph(style=style)
    add_markdown_runs(paragraph, text)
    return paragraph


def is_table_separator(line: str) -> bool:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in cells)


def collect_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        if not is_table_separator(lines[index]):
            rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
        index += 1
    return rows, index


def choose_widths(rows: list[list[str]]) -> list[float]:
    columns = max(len(row) for row in rows)
    if columns == 2:
        return [1.75, 4.75]
    if columns == 3:
        return [1.45, 2.35, 2.70]
    if columns == 4:
        return [1.25, 1.65, 1.80, 1.80]
    return [6.5 / columns] * columns


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    widths = choose_widths(rows)
    for row_index, row in enumerate(rows):
        for col_index in range(columns):
            cell = table.cell(row_index, col_index)
            text = row[col_index] if col_index < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_markdown_runs(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
            if row_index == 0:
                set_cell_shading(cell, "E8EEF5")
                for run in paragraph.runs:
                    run.bold = True
            if col_index == 0 and columns > 2:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_geometry(table, widths)
    doc.add_paragraph()


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def add_cover(doc: Document) -> None:
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(6)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("修改后新版功能、技术点与旧版对比文档")
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("0B2545")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle.add_run("agent_app_v2 技术说明与面试复习版").italic = True

    meta = doc.add_table(rows=4, cols=2)
    meta.style = "Table Grid"
    items = [
        ("整理日期", "2026-07-26"),
        ("新版项目", r"C:\Users\HP\vsocde\app_v2\agent_app_v2"),
        ("旧版项目", r"C:\Users\HP\Desktop\app.zip"),
        ("文档用途", "功能说明、技术点复习、新旧版本对比、面试问答准备"),
    ]
    for row, (label, value) in zip(meta.rows, items, strict=True):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], "E8EEF5")
        for cell in row.cells:
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    if cell is row.cells[0]:
                        run.bold = True
    set_table_geometry(meta, [1.35, 5.15])
    doc.add_paragraph()


def add_code_block(doc: Document, code_lines: list[str]) -> None:
    for line in code_lines:
        paragraph = doc.add_paragraph(style="Normal")
        paragraph.paragraph_format.left_indent = Inches(0.2)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(line if line else " ")
        style_run(run, code=True)
    doc.add_paragraph()


def build(output: Path = DEFAULT_OUTPUT) -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    in_code = False
    code_lines: list[str] = []
    index = 0
    while index < len(lines):
        raw = lines[index]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        if (
            stripped.startswith("|")
            and index + 1 < len(lines)
            and is_table_separator(lines[index + 1])
        ):
            rows, next_index = collect_table(lines, index)
            add_table(doc, rows)
            index = next_index
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            if level == 1:
                index += 1
                continue
            add_paragraph(doc, text, style=f"Heading {min(level - 1, 3)}")
            index += 1
            continue

        number_match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if number_match:
            add_paragraph(doc, number_match.group(2), style="List Number")
            index += 1
            continue

        bullet_match = re.match(r"^[-*]\s+(.*)$", stripped)
        if bullet_match:
            add_paragraph(doc, bullet_match.group(1), style="List Bullet")
            index += 1
            continue

        quote_match = re.match(r"^>\s+(.*)$", stripped)
        if quote_match:
            paragraph = add_paragraph(doc, quote_match.group(1), style="Normal")
            paragraph.paragraph_format.left_indent = Inches(0.25)
            for run in paragraph.runs:
                run.italic = True
                run.font.color.rgb = RGBColor.from_string("1F4D78")
            index += 1
            continue

        add_paragraph(doc, stripped)
        index += 1

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("agent_app_v2 技术对比文档")

    doc.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build the comparison DOCX from Markdown")
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    args = parser.parse_args()
    build(Path(args.output))
