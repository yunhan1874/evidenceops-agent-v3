#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "EvidenceOps_Agent_升级对比与大厂面试手册.md"
DEFAULT_OUTPUT = ROOT / "EvidenceOps_Agent_v3_与最早版app_升级对比与大厂面试手册.docx"

FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft YaHei"
NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5E6B78"
GRID = "B8C4D0"
HEADER_FILL = "E8EEF5"
CALLOUT_FILL = "F4F6F9"
CODE_FILL = "F6F8FA"


def set_font(run, *, name=FONT_LATIN, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top=80, bottom=80, start=120, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        element = tc_mar.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_table_geometry(table, widths_in: list[float]) -> None:
    widths_dxa = [round(width * 1440) for width in widths_in]
    total_dxa = sum(widths_dxa)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, width in enumerate(widths_dxa):
            cell = row.cells[index]
            cell.width = Inches(widths_in[index])
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, value, end):
        run._r.append(element)
    set_font(run, size=9, color=MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    heading_tokens = (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 11.5, DARK_BLUE, 10, 5),
    )
    for name, size, color, before, after in heading_tokens:
        style = doc.styles[name]
        style.font.name = FONT_LATIN
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = FONT_LATIN
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("最早版 app → EvidenceOps Agent v3  |  升级对比与大厂面试手册")
    set_font(run, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    prefix = footer.add_run("项目技术手册  ·  ")
    set_font(prefix, size=9, color=MUTED)
    add_page_field(footer)


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(42)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(14)
    run = kicker.add_run("PROJECT DEEP-DIVE  /  INTERVIEW EDITION")
    set_font(run, size=10, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("EvidenceOps Agent v3")
    set_font(run, size=29, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("从最早版 app.zip 到可信 EvidenceOps Agent")
    set_font(run, size=15, color=DARK_BLUE)

    lead = doc.add_paragraph()
    lead.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead.paragraph_format.left_indent = Inches(0.45)
    lead.paragraph_format.right_indent = Inches(0.45)
    lead.paragraph_format.space_after = Pt(22)
    run = lead.add_run(
        "从 2617 行单文件 RAG 代理，升级为具备状态图、证据门禁、"
        "人工审批、幂等、CAS 和对账语义的可信 Agent。"
    )
    set_font(run, size=11.5, italic=True, color=MUTED)

    metrics = doc.add_table(rows=2, cols=4)
    labels = ("单元测试", "端到端用例", "质量指标", "并发审批")
    values = ("26 / 26", "200 / 200", "9 项 1.0", "200 + 409 / 1 次反馈")
    for index, label in enumerate(labels):
        cell = metrics.cell(0, index)
        set_cell_shading(cell, HEADER_FILL)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(label)
        set_font(run, size=9, bold=True, color=DARK_BLUE)
    for index, value in enumerate(values):
        cell = metrics.cell(1, index)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(value)
        set_font(run, size=9.5, bold=True, color=NAVY)
    set_repeat_table_header(metrics.rows[0])
    set_table_geometry(metrics, [1.625, 1.625, 1.625, 1.625])

    doc.add_paragraph()
    rows = (
        ("升级前项目", "用户原始 app.zip（app/app.py）"),
        ("升级后项目", "agent_app_evidence_v3"),
        ("技术栈", "FastAPI · LangGraph · PostgreSQL · SQLAlchemy Async · RAGFlow"),
        ("用途", "秋招项目讲解 · 技术复习 · 深挖问答 · 简历表达"),
    )
    for label, value in rows:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.55)
        paragraph.paragraph_format.right_indent = Inches(0.35)
        paragraph.paragraph_format.space_after = Pt(3)
        label_run = paragraph.add_run(f"{label}：")
        set_font(label_run, size=9.5, bold=True, color=NAVY)
        value_run = paragraph.add_run(value)
        set_font(value_run, size=9.5)

    navigation = doc.add_paragraph()
    navigation.paragraph_format.space_before = Pt(18)
    navigation.paragraph_format.space_after = Pt(0)
    navigation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = navigation.add_run("阅读路径：开场陈述 → 最早版源码审计 → 架构演进 → 故障矩阵 → 69 道面试题")
    set_font(run, size=9.5, color=MUTED)
    doc.add_page_break()


def add_markdown_runs(paragraph, text: str) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, name="Consolas", size=9.2, color=DARK_BLUE)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, bold=True, color=NAVY)
        else:
            run = paragraph.add_run(part)
            set_font(run)


def add_paragraph(doc: Document, text: str, *, style="Normal"):
    paragraph = doc.add_paragraph(style=style)
    add_markdown_runs(paragraph, text)
    return paragraph


def add_manual_numbered_paragraph(doc: Document, number: str, text: str) -> None:
    """Preserve the Markdown list number so separate lists restart correctly."""
    paragraph = doc.add_paragraph(style="List Paragraph")
    paragraph.paragraph_format.left_indent = Inches(0.34)
    paragraph.paragraph_format.first_line_indent = Inches(-0.22)
    add_markdown_runs(paragraph, f"{number}. {text}")


def add_quote(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.15)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.2
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), CALLOUT_FILL)
    p_pr.append(shading)
    add_markdown_runs(paragraph, text)
    for run in paragraph.runs:
        run.italic = True
        run.font.color.rgb = RGBColor.from_string(DARK_BLUE)


def add_code_block(doc: Document, lines: list[str]) -> None:
    for index, line in enumerate(lines):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.right_indent = Inches(0.12)
        paragraph.paragraph_format.space_before = Pt(4 if index == 0 else 0)
        paragraph.paragraph_format.space_after = Pt(4 if index == len(lines) - 1 else 0)
        paragraph.paragraph_format.line_spacing = 1.0
        p_pr = paragraph._p.get_or_add_pPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), CODE_FILL)
        p_pr.append(shading)
        run = paragraph.add_run(line if line else " ")
        set_font(run, name="Consolas", size=8.3, color=NAVY)


def is_table_separator(line: str) -> bool:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in cells)


def collect_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        if not is_table_separator(lines[index]):
            rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
        index += 1
    return rows, index


def choose_widths(column_count: int) -> list[float]:
    patterns = {
        2: [1.7, 4.8],
        3: [1.4, 2.55, 2.55],
        4: [1.15, 1.65, 2.15, 1.55],
    }
    return patterns.get(column_count, [6.5 / column_count] * column_count)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    font_size = 8.4 if columns >= 4 else 9.2
    for row_index, row in enumerate(rows):
        for column_index in range(columns):
            cell = table.cell(row_index, column_index)
            text = row[column_index] if column_index < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            if column_index == 0 and columns >= 3:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_markdown_runs(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(font_size)
            if row_index == 0:
                set_cell_shading(cell, HEADER_FILL)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(NAVY)
    set_table_geometry(table, choose_widths(columns))
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def build(output: Path) -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    code_lines: list[str] = []
    in_code = False
    index = 0
    force_break_before = {"大厂面试题与参考答案", "简历描述建议", "最后一分钟复习清单"}
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
            in_code = not in_code
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not stripped:
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            rows, index = collect_table(lines, index)
            add_table(doc, rows)
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                index += 1
                continue
            heading_key = re.sub(r"^\d+(?:\.\d+)*\.\s*", "", text)
            if heading_key in force_break_before:
                doc.add_page_break()
            style = {2: "Heading 1", 3: "Heading 2", 4: "Heading 3"}[level]
            paragraph = add_paragraph(doc, text, style=style)
            if level == 4:
                paragraph.paragraph_format.keep_with_next = True
            index += 1
            continue

        quote = re.match(r"^>\s?(.*)$", stripped)
        if quote:
            add_quote(doc, quote.group(1))
            index += 1
            continue

        numbered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if numbered:
            add_manual_numbered_paragraph(doc, numbered.group(1), numbered.group(2))
            index += 1
            continue

        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        if bullet:
            add_paragraph(doc, bullet.group(1), style="List Bullet")
            index += 1
            continue

        add_paragraph(doc, stripped)
        index += 1

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "最早版 app 到 EvidenceOps Agent v3 升级对比与大厂面试手册"
    doc.core_properties.subject = "最早版 app.zip 源码审计、企业知识 Agent 升级对比、技术点与面试问答"
    doc.core_properties.keywords = "FastAPI, RAGFlow, LangGraph, Evidence Guard, 异步IO, 幂等, 并发"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build the EvidenceOps Agent interview handbook")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    arguments = parser.parse_args()
    build(arguments.output)
