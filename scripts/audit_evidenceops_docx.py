#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def nearly_inches(value, expected: float, tolerance: float = 0.01) -> bool:
    return value is not None and abs(value.inches - expected) <= tolerance


def main() -> None:
    parser = argparse.ArgumentParser(description="Structural audit for the EvidenceOps handbook")
    parser.add_argument("docx", type=Path)
    args = parser.parse_args()

    doc = Document(args.docx)
    paragraphs = list(doc.paragraphs)
    question_numbers = []
    leaked_markdown = []
    for index, paragraph in enumerate(paragraphs, start=1):
        text = paragraph.text.strip()
        if paragraph.style.name == "Heading 3":
            match = re.match(r"Q(\d+)：", text)
            if match:
                question_numbers.append(int(match.group(1)))
        if re.match(r"^(#{1,4}\s|```|\|\s*---|\*\*.+\*\*$)", text):
            leaked_markdown.append({"paragraph": index, "text": text[:80]})

    section = doc.sections[0]
    page_geometry_ok = all(
        (
            nearly_inches(section.page_width, 8.5),
            nearly_inches(section.page_height, 11.0),
            nearly_inches(section.left_margin, 1.0),
            nearly_inches(section.right_margin, 1.0),
            nearly_inches(section.top_margin, 1.0),
            nearly_inches(section.bottom_margin, 1.0),
        )
    )
    header_rows = 0
    fixed_height_rows = 0
    for table in doc.tables:
        if table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None:
            header_rows += 1
        for row in table.rows:
            if row._tr.get_or_add_trPr().find(qn("w:trHeight")) is not None:
                fixed_height_rows += 1

    result = {
        "file": str(args.docx),
        "paragraph_count": len(paragraphs),
        "table_count": len(doc.tables),
        "question_count": len(question_numbers),
        "question_sequence_ok": question_numbers == list(range(1, 70)),
        "markdown_leaks": leaked_markdown,
        "page_geometry_ok": page_geometry_ok,
        "all_tables_have_header_rows": header_rows == len(doc.tables),
        "fixed_height_row_count": fixed_height_rows,
        "footer_has_page_field": "PAGE" in section.footer._element.xml,
        "core_author_blank": not bool(doc.core_properties.author),
    }
    failures = [
        not result["question_sequence_ok"],
        bool(result["markdown_leaks"]),
        not result["page_geometry_ok"],
        not result["all_tables_have_header_rows"],
        result["fixed_height_row_count"] != 0,
        not result["footer_has_page_field"],
        not result["core_author_blank"],
    ]
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if any(failures):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
