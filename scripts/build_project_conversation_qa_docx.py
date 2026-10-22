#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
from pathlib import Path

from build_evidenceops_interview_docx import (
    FONT_CJK,
    FONT_LATIN,
    add_code_block,
    add_manual_numbered_paragraph,
    add_markdown_runs,
    add_page_field,
    add_paragraph,
    add_quote,
    add_table,
    collect_table,
    configure_document,
    is_table_separator,
    set_font,
)
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "PROJECT_CONVERSATION_QA.md"
DEFAULT_OUTPUT = ROOT / "EvidenceOps_Agent_项目迭代问答整理.docx"


def configure_qa_document(doc: Document) -> None:
    configure_document(doc)
    black = RGBColor(0, 0, 0)
    for style_name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = FONT_LATIN
        style._element.get_or_add_rPr().rFonts.set(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
            FONT_CJK,
        )
        style.font.color.rgb = black

    title_properties = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    header = doc.sections[0].header.paragraphs[0]
    header.clear()
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("EvidenceOps Agent 项目迭代问答整理"), size=9, color="666666")

    footer = doc.sections[0].footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run("秋招项目复盘  "), size=9, color="666666")
    add_page_field(footer)


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(70)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run("EvidenceOps Agent 项目迭代问答整理")
    set_font(run, size=26, bold=True, color="000000")
    title_border = title._p.get_or_add_pPr().find(qn("w:pBdr"))
    if title_border is not None:
        title._p.get_or_add_pPr().remove(title_border)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(34)
    run = subtitle.add_run("秋招项目复盘与面试表达")
    set_font(run, size=14, color="000000")

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.left_indent = Inches(0.7)
    intro.paragraph_format.right_indent = Inches(0.7)
    intro.paragraph_format.space_after = Pt(28)
    run = intro.add_run(
        "从最早版 RAGFlow 后端到具备证据门禁 人工审批 幂等并发控制与 Execution Harness 的企业知识 Agent"
    )
    set_font(run, size=11.5, color="444444")

    facts = doc.add_table(rows=5, cols=2)
    rows = (
        ("比较基线", "用户最早提供的 app.zip"),
        ("升级目录", "app_v2/agent_app_evidence_v3"),
        ("当前自动化验证", "36 个单元测试通过 静态检查通过"),
        ("评测资产", "90 条黄金集 70 条对抗集 40 条回放种子"),
        ("用途", "项目复盘 简历修改 面试深挖"),
    )
    for row_index, (label, value) in enumerate(rows):
        facts.cell(row_index, 0).text = label
        facts.cell(row_index, 1).text = value
        facts.cell(row_index, 0).paragraphs[0].runs[0].bold = True
    facts.autofit = False
    for row in facts.rows:
        row.cells[0].width = Inches(1.5)
        row.cells[1].width = Inches(5.0)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(24)
    set_font(
        note.add_run("事实口径 真实 Qwen 与 200 条全量结果为历史版本报告 最新前置门禁另行完成单元和 Docker 端到端验证"),
        size=9.5,
        color="555555",
    )
    doc.add_page_break()


def build(output: Path) -> None:
    doc = Document()
    configure_qa_document(doc)
    add_cover(doc)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_lines: list[str] = []
    index = 0
    force_break_before = {
        "三 Execution Harness 与最新面试考点",
        "四 测试集与完整测试逻辑",
    }

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
            in_code = not in_code
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            rows, index = collect_table(lines, index)
            add_table(doc, rows)
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                index += 1
                continue
            if text in force_break_before:
                doc.add_page_break()
            style = {2: "Heading 1", 3: "Heading 2", 4: "Heading 3"}[level]
            paragraph = add_paragraph(doc, text, style=style)
            paragraph.runs[0].font.color.rgb = RGBColor(0, 0, 0)
            index += 1
            continue
        quote = re.match(r"^>\s?(.*)$", stripped)
        if quote:
            add_quote(doc, quote.group(1))
            index += 1
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if numbered:
            add_manual_numbered_paragraph(doc, numbered.group(1), numbered.group(2))
            index += 1
            continue
        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        if bullet:
            add_paragraph(doc, bullet.group(1), style="List Bullet")
            index += 1
            continue
        paragraph = doc.add_paragraph()
        add_markdown_runs(paragraph, stripped)

        index += 1

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "EvidenceOps Agent 项目迭代问答整理"
    doc.core_properties.subject = "企业知识 Agent 项目升级 Harness 测试 Qwen 接入与简历表达"
    doc.core_properties.keywords = "FastAPI LangGraph RAGFlow PostgreSQL Qwen Harness Evidence Guard"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build project conversation Q and A DOCX")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    arguments = parser.parse_args()
    build(arguments.output)
